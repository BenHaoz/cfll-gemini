"""知网（CNKI）导出文件导入：补文献中心未收录的核心期刊（如《广西民族研究》）。
支持的导出格式：RefWorks、EndNote、NoteExpress（文本），以及“自定义”导出的 Excel/CSV 表格。"""
from __future__ import annotations

import csv
import io
import re
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

import yaml

from .config import CONFIG_DIR
from .pubs import (CORE_TIERS, PUBS_DIR, Article, InstitutionMatcher, article_key, issue_to_month, load_articles, norm,
                   save_articles, split_author_field, upsert)

CNKI_DIR = PUBS_DIR / "cnki"
IMPORT_SUFFIXES = {".txt", ".ris", ".xls", ".xlsx", ".csv", ".html", ".htm"}

REFWORKS_TAGS = {"T1": "title", "A1": "authors", "AD": "organs", "JF": "journal", "YR": "year", "IS": "issue", "LK": "url", "UL": "url"}
ENDNOTE_TAGS = {"%T": "title", "%A": "authors", "%+": "organs", "%J": "journal", "%D": "year", "%N": "issue", "%U": "url"}
NOTEEXPRESS_TAGS = {"Title": "title", "Author": "authors", "Author Address": "organs", "Journal": "journal", "Year": "year",
                    "Issue": "issue", "URL": "url"}
TABLE_HEADERS = {
    "title": {"title", "题名", "篇名", "标题"},
    "authors": {"author", "作者"},
    "organs": {"organ", "单位", "机构", "作者单位", "author address"},
    "journal": {"source", "文献来源", "刊名", "来源", "journal"},
    "year": {"year", "年", "年份"},
    "issue": {"period", "期", "issue"},
    "pubtime": {"pubtime", "发表时间", "出版日期"},
    "url": {"url", "链接"},
}


def _record(fields: dict[str, list[str]]) -> dict[str, Any]:
    def first(k: str) -> str:
        return (fields.get(k) or [""])[0].strip()
    authors: list[str] = []
    for v in fields.get("authors", []):
        authors += [a for a in split_author_field(v) if a not in authors]
    organs: list[str] = []
    for v in fields.get("organs", []):
        organs += [o.strip() for o in re.split(r"[;；]", v) if o.strip() and o.strip() not in organs]
    pubtime = first("pubtime")
    m_year = re.search(r"(19|20)\d{2}", first("year") or pubtime)
    m_issue = re.fullmatch(r"0*(\d{1,2})", first("issue"))
    m_date = re.match(r"((?:19|20)\d{2})[-/.年](\d{1,2})", pubtime)
    return {"title": re.sub(r"\s+", " ", first("title")), "authors": authors, "organs": organs[:10], "journal": first("journal"),
            "year": int(m_year.group(0)) if m_year else 0, "issue": int(m_issue.group(1)) if m_issue else 0,
            "date": f"{m_date.group(1)}-{int(m_date.group(2)):02d}" if m_date else "", "url": first("url")}


def _parse_tagged(text: str, tag_re: str, tags: dict[str, str], start_tag: str) -> list[dict[str, Any]]:
    """逐行标签格式（RefWorks/EndNote/NoteExpress）：遇到起始标签开新记录。"""
    recs: list[dict[str, list[str]]] = []
    cur: dict[str, list[str]] | None = None
    last: str | None = None
    for line in text.splitlines():
        m = re.match(tag_re, line)
        if m:
            tag, val = m.group(1), m.group(2).strip()
            if tag == start_tag:
                cur = {}
                recs.append(cur)
            last = tags.get(tag)
            if cur is not None and last:
                cur.setdefault(last, []).append(val)
        elif cur is not None and last in ("title", "organs") and line.strip():
            cur[last][-1] += line.strip()
    return [_record(r) for r in recs]


def _header_field(h: str) -> str | None:
    parts = {p.strip().lower() for p in re.split(r"[-－]", str(h or "")) if p.strip()}
    for field, names in TABLE_HEADERS.items():
        if parts & names:
            return field
    return None


def _parse_rows(rows: Iterable[list[Any]]) -> list[dict[str, Any]]:
    rows = [[("" if c is None else str(c)).strip() for c in r] for r in rows]
    for i, r in enumerate(rows):
        cols = {j: _header_field(h) for j, h in enumerate(r)}
        found = {f for f in cols.values() if f}
        if {"title", "journal"} <= found:
            out = []
            for data in rows[i + 1:]:
                fields: dict[str, list[str]] = {}
                for j, f in cols.items():
                    if f and j < len(data) and data[j]:
                        fields.setdefault(f, []).append(data[j])
                if fields.get("title"):
                    out.append(_record(fields))
            return out
    raise ValueError("表格中找不到“题名/Title”与“文献来源/Source”列")


def _parse_html_table(text: str) -> list[dict[str, Any]]:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(text, "lxml")
    rows = [[c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"])] for tr in soup.find_all("tr")]
    return _parse_rows(rows)


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def parse_cnki_file(path: Path) -> list[dict[str, Any]]:
    raw = Path(path).read_bytes()
    if raw[:2] == b"PK":
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
        return _parse_rows(ws.iter_rows(values_only=True))
    if raw[:4] == b"\xd0\xcf\x11\xe0":
        raise ValueError(f"{Path(path).name} 是旧版 .xls 二进制格式：请另存为 .xlsx，或在知网改用 RefWorks / EndNote 格式导出")
    text = _decode(raw)
    if "<table" in text[:20000].lower():
        return _parse_html_table(text)
    if re.search(r"^RT\s", text, re.M):
        return _parse_tagged(text, r"^([A-Z][A-Z0-9])\s+(.*)$", REFWORKS_TAGS | {"RT": "_start"}, "RT")
    if re.search(r"^%0\s", text, re.M):
        return _parse_tagged(text, r"^(%.)\s+(.*)$", ENDNOTE_TAGS | {"%0": "_start"}, "%0")
    if "{Reference Type}" in text:
        return _parse_tagged(text, r"^\{([^}]+)\}:\s*(.*)$", NOTEEXPRESS_TAGS | {"Reference Type": "_start"}, "Reference Type")
    return _parse_rows(csv.reader(io.StringIO(text)))


def import_records(records: list[dict[str, Any]], arts: dict[str, Article], core_journals: list[dict[str, Any]],
                   matcher: InstitutionMatcher) -> dict[str, Any]:
    """只收监测名单中等级为核心的期刊（与文献中心采集口径一致）；其余计入 skipped 并说明原因。"""
    by_name: dict[str, dict[str, Any]] = {}
    for j in core_journals:
        for n in [j["name"], *(j.get("aliases") or [])]:
            by_name[norm(n)] = j
    new: list[Article] = []
    skipped: Counter = Counter()
    for r in records:
        j = by_name.get(norm(r["journal"]))
        if not j:
            skipped[f"{r['journal'] or '(无刊名)'}：不在监测名单"] += 1
            continue
        if j.get("tier") not in CORE_TIERS:
            skipped[f"{j['name']}：等级为“{j.get('tier')}”，不计入统计"] += 1
            continue
        if not r["title"] or not r["year"]:
            skipped["缺题名或年份"] += 1
            continue
        freq = j.get("frequency", "双月刊")
        date = r["date"] or (f"{r['year']}-{issue_to_month(r['issue'], freq):02d}" if r["issue"] else "")
        new.append(Article(key=article_key(j["name"], r["year"], r["issue"], r["title"]), title=r["title"], journal=j["name"],
                           year=r["year"], issue=r["issue"], date=date, authors=r["authors"], affiliations=r["organs"],
                           institutions=matcher.match_all(r["organs"]), url=r["url"], tier=j.get("tier", ""), source="cnki"))
    added = upsert(arts, new)
    return {"parsed": len(records), "eligible": len(new), "added": added, "merged": len(new) - added,
            "skipped": dict(skipped.most_common())}


def import_cnki(paths: list[Path] | None = None) -> dict[str, Any]:
    paths = paths or sorted(p for p in CNKI_DIR.glob("*") if p.suffix.lower() in IMPORT_SUFFIXES)
    core = yaml.safe_load(open(CONFIG_DIR / "journals_core.yaml", encoding="utf-8")) or {}
    insts = yaml.safe_load(open(CONFIG_DIR / "institutions.yaml", encoding="utf-8")) or {}
    matcher = InstitutionMatcher(insts.get("institutions", []))
    records: list[dict[str, Any]] = []
    errors: dict[str, str] = {}
    for p in paths:
        try:
            records += parse_cnki_file(p)
        except ValueError as exc:
            errors[p.name] = str(exc)
    arts = load_articles()
    res = import_records(records, arts, core.get("journals", []), matcher)
    if res["eligible"]:
        save_articles(arts)
    res["files"] = [p.name for p in paths]
    res["errors"] = errors
    return res
